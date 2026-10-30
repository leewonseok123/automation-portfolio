from datetime import date, datetime

import excel_report as er
import make_sample_data
import pandas as pd
import pytest
from openpyxl import load_workbook


@pytest.mark.parametrize("raw", ["2026-09-03", "2026.09.03", "20260903", datetime(2026, 9, 3), date(2026, 9, 3)])
def test_parse_date_formats(raw):
    assert er.parse_date(raw) == date(2026, 9, 3)


@pytest.mark.parametrize("raw", ["미정", "", "2026-13-45"])
def test_parse_date_rejects_garbage(raw):
    assert er.parse_date(raw) is None


@pytest.mark.parametrize("raw,expected", [("12,000원", 12000), ("₩12,000", 12000), (12000, 12000), ("-3", -3), ("", None)])
def test_parse_number(raw, expected):
    assert er.parse_number(raw) == expected


def test_map_columns_aliases_and_missing():
    assert er.map_columns(["판매일", "메뉴", "qty", "price", "total"]) == {
        "판매일": "date",
        "메뉴": "product",
        "qty": "qty",
        "price": "price",
        "total": "amount",
    }
    with pytest.raises(ValueError, match="amount"):
        er.map_columns(["일자", "상품명", "수량", "단가"])


def test_locate_header_skips_title_rows():
    raw = pd.DataFrame([["성수점 매출", None, None, None, None], [None] * 5, ["일자", "상품명", "수량", "단가", "금액"]])
    assert er.locate_header(raw) == 2


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    base = tmp_path_factory.mktemp("excel")
    make_sample_data.main(base / "input")
    files = sorted((base / "input").iterdir())
    issues: list[dict] = []
    df = pd.concat([er.load_branch(f, issues) for f in files], ignore_index=True)
    summary = er.build_report(df, issues, files, base / "report.xlsx")
    return df, issues, summary, base / "report.xlsx"


def test_all_branches_loaded(report):
    df, _, summary, _ = report
    assert set(df["branch"]) == {"성수점", "강남점", "홍대점", "판교점", "부산점"}
    assert len(summary["files"]) == 5


def test_every_planted_problem_is_caught(report):
    _, _, summary, _ = report
    for kind in ("빈 행 제거", "중복 행 제거", "금액 불일치", "형식 오류로 제외", "반품"):
        assert summary["issues_by_kind"].get(kind, 0) > 0, kind


def test_amounts_are_consistent_after_cleaning(report):
    df, _, summary, _ = report
    assert ((df["qty"] * df["price"] - df["amount"]).abs() <= 1).all()
    assert summary["total_revenue"] == int(df["amount"].sum())
    assert not df.duplicated(subset=["branch", "date", "product", "qty", "price", "amount"]).any()


def test_report_workbook_structure(report):
    *_, path = report
    wb = load_workbook(path)
    assert wb.sheetnames == ["요약", "지점×월", "상품 순위", "정제 데이터", "검증 로그"]
    assert len(wb["요약"]._charts) == 2
    assert len(wb["상품 순위"]._charts) == 1
