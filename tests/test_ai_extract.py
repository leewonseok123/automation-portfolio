import json
import re
from types import SimpleNamespace

import ai_extract as ax
import pytest

INPUT = ax.HERE / "input"
TRUTH = json.loads((INPUT / "ground_truth.json").read_text(encoding="utf-8"))
FILES = sorted(INPUT.glob("*.pdf"))


class TruthExtractor:
    """정답을 그대로 돌려주는 가짜 AI — 파이프라인(검증·평가)만 따로 시험합니다."""

    name, model = "fake", "truth"

    def __init__(self, truth: dict, mutate=None):
        self.truth, self.mutate, self.calls = truth, mutate, 0

    def extract(self, text):
        self.calls += 1
        doc = json.loads(json.dumps(self.truth))
        if self.mutate:
            self.mutate(doc)
        return doc, ax.Usage(100, 50)


def run_all(mutate_file=None, mutate=None):
    results = []
    for f in FILES:
        m = mutate if f.name == mutate_file else None
        results.append(ax.process(f, TruthExtractor(TRUTH[f.name], m)))
    return results


def codes(result):
    return {i.code for i in result.issues}


def by_name(results):
    return {r.file: r for r in results}


def test_sample_set_is_complete():
    assert len(FILES) == 14
    assert set(TRUTH) == {f.name for f in FILES}


def test_biz_no_checksum():
    assert ax.biz_no_valid(TRUTH["doc01_tax_invoice.pdf"]["supplier"]["biz_no"])
    assert not ax.biz_no_valid(TRUTH["doc04_tax_invoice.pdf"]["supplier"]["biz_no"])
    assert not ax.biz_no_valid("123-45-678")
    assert not ax.biz_no_valid(None)


def test_planted_flaws_are_caught_and_clean_docs_pass():
    r = by_name(run_all())
    assert "TOTAL_SUM" in codes(r["doc03_tax_invoice.pdf"])
    assert "BIZNO_CHECKSUM" in codes(r["doc04_tax_invoice.pdf"])
    assert "ITEMS_SUM" in codes(r["doc05_tax_invoice.pdf"])
    for name, t in TRUTH.items():
        if t["planted_flaw"] is None and t["doc_type"] != "other" and not t.get("scanned"):
            assert r[name].status == "ok", (name, r[name].issues)


def test_non_invoice_is_skipped_and_scan_never_calls_the_api():
    ex = TruthExtractor(TRUTH["doc14_scan.pdf"])
    scan = ax.process(INPUT / "doc14_scan.pdf", ex)
    assert scan.status == "needs_ocr" and ex.calls == 0
    assert by_name(run_all())["doc13_other.pdf"].status == "skipped"


def test_evaluate_perfect_when_extraction_matches_truth():
    ev = ax.evaluate(run_all(), TRUTH)
    assert ev["field_accuracy"] == 100.0
    assert ev["docs_scored"] == 13  # 스캔본 제외
    assert ev["docs_perfect"] == 13


def test_evaluate_detects_a_wrong_value():
    ev = ax.evaluate(run_all("doc07_quotation.pdf", lambda d: d.update(total=d["total"] + 1)), TRUTH)
    assert ev["field_accuracy"] < 100
    assert ev["wrong_fields"]["doc07_quotation.pdf"] == ["total"]
    assert ev["docs_perfect"] == 12


def test_rules_baseline_is_clearly_weaker():
    results = [ax.process(f, ax.RulesExtractor()) for f in FILES]
    assert ax.evaluate(results, TRUTH)["field_accuracy"] < 60


@pytest.mark.parametrize("raw", ['```json\n{"a": 1}\n```', '결과입니다:\n{"a": 1}\n끝', '{"a": 1}'])
def test_parse_json_tolerates_wrappers(raw):
    assert ax.parse_json(raw) == {"a": 1}


@pytest.mark.parametrize("raw", ["", "no json here", "[1, 2]"])
def test_parse_json_rejects_garbage(raw):
    with pytest.raises(ValueError):
        ax.parse_json(raw)


def test_normalize_handles_messy_model_output():
    doc = ax.normalize(
        {
            "doc_type": "Tax_Invoice",
            "issue_date": "2026년 9월 3일",
            "supplier": {"name": " 다온푸드 ", "biz_no": "9698797302"},
            "items": [{"name": "의자", "qty": "5", "unit_price": "₩145,000", "amount": "725,000원"}],
            "supply_amount": "725,000원",
            "vat": 72500.0,
            "total": "₩797,500",
        }
    )
    assert doc["doc_type"] == "tax_invoice"
    assert doc["issue_date"] == "2026-09-03"
    assert doc["supplier"] == {"name": "다온푸드", "biz_no": "969-87-97302"}
    assert doc["items"][0] == {"name": "의자", "qty": 5, "unit_price": 145000, "amount": 725000}
    assert (doc["supply_amount"], doc["vat"], doc["total"]) == (725000, 72500, 797500)


@pytest.mark.parametrize(
    "raw,iso", [("2026.09.03", "2026-09-03"), ("26-09-03", "2026-09-03"), ("2026-02-30", None), ("미정", None)]
)
def test_date_formats(raw, iso):
    assert ax.to_iso_date(raw) == iso


def test_unknown_doc_type_becomes_other():
    assert ax.normalize({"doc_type": "invoice-ish"})["doc_type"] == "other"


class FakeClient:
    def __init__(self, replies):
        self.replies, self.calls = list(replies), 0
        self.chat = SimpleNamespace(completions=self)

    def create(self, **_kwargs):
        self.calls += 1
        content = self.replies.pop(0)
        return SimpleNamespace(
            choices=[SimpleNamespace(message=SimpleNamespace(content=content))],
            usage=SimpleNamespace(prompt_tokens=100, completion_tokens=20),
        )


def make_deepseek(client, attempts=3):
    ex = ax.DeepSeekExtractor.__new__(ax.DeepSeekExtractor)
    ex.client, ex.model, ex.attempts = client, "test-model", attempts
    return ex


def test_deepseek_retries_after_malformed_json_and_sums_usage():
    client = FakeClient(["죄송합니다, 다음과 같습니다", '{"doc_type": "other"}'])
    raw, usage = make_deepseek(client).extract("문서 텍스트")
    assert raw == {"doc_type": "other"}
    assert client.calls == 2 and (usage.prompt, usage.completion) == (200, 40)


def test_deepseek_gives_up_and_process_reports_failure():
    ex = make_deepseek(FakeClient(["x", "y", "z"]))
    res = ax.process(INPUT / "doc01_tax_invoice.pdf", ex)
    assert res.status == "failed" and "JSON" in res.error


def test_missing_api_key_exits_with_clear_message(monkeypatch):
    monkeypatch.delenv("DEEPSEEK_API_KEY", raising=False)
    with pytest.raises(SystemExit, match="DEEPSEEK_API_KEY"):
        ax.DeepSeekExtractor()


def test_prompt_forbids_guessing_and_demands_json():
    assert "계산해서 채우지" in ax.SYSTEM_PROMPT or "계산해서" in ax.SYSTEM_PROMPT
    assert re.search(r"JSON", ax.SYSTEM_PROMPT)


def test_excel_output_has_expected_sheets(tmp_path):
    results = run_all()
    summary = {
        "provider": "fake",
        "model": "truth",
        "docs": len(results),
        "status_counts": {"ok": 1},
        "tokens": {"prompt": 1, "completion": 1},
        "accuracy": None,
    }
    out = tmp_path / "x.xlsx"
    ax.write_xlsx(results, summary, out)
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert wb.sheetnames == ["추출결과", "품목", "검토 필요", "요약"]
    assert wb["추출결과"].max_row == len(results) + 1
    assert wb["검토 필요"].max_row > 1


def test_sample_generator_is_reproducible(tmp_path):
    pytest.importorskip("reportlab")
    import make_sample_docs as gen

    try:
        gen.register_font()
    except SystemExit:
        pytest.skip("한글 폰트 없음")
    gen.main(tmp_path)
    assert json.loads((tmp_path / "ground_truth.json").read_text(encoding="utf-8")) == TRUTH
