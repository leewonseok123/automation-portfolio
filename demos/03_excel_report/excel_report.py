"""
03 · Excel Report — 양식이 제각각인 지점별 매출 파일 → 통합 보고서 자동 생성

기능
- input/ 폴더의 xlsx·csv 를 모두 읽음 (CSV 인코딩 자동 감지: UTF-8 → CP949)
- 표 위 제목 행이 있어도 헤더 행을 자동 탐지
- 컬럼명 별칭 매핑 (일자/날짜/판매일 → date …)
- 날짜·금액 형식 정규화, 공백 제거, 빈 행·중복 행 제거
- 수량×단가 ≠ 금액 같은 이상 데이터는 '검증 로그' 시트로 분리
- 요약 KPI · 지점×월 피벗 · 상품 순위 · 요일 패턴 + 엑셀 네이티브 차트

사용법
    python make_sample_data.py      # 예시 입력 파일 생성
    python excel_report.py          # output/매출_통합보고서.xlsx 생성
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import time
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).parent
log = logging.getLogger("excel_report")

ALIASES = {
    "date": ["일자", "날짜", "판매일", "거래일자", "date"],
    "product": ["상품명", "품목", "메뉴", "상품", "product", "item"],
    "qty": ["수량", "판매수량", "qty", "quantity"],
    "price": ["단가", "판매가", "price", "unit_price"],
    "amount": ["금액", "매출액", "합계", "total", "amount"],
}
WEEKDAYS = ["월", "화", "수", "목", "금", "토", "일"]

INK, BONE, PAPER, ASH, VERMILION = "131311", "F2EFE9", "E9E5DD", "6F6C66", "E8442E"
THIN = Side(style="thin", color="D5D0C6")


# ───────────────────────────── 읽기 · 정규화 ─────────────────────────────

def read_raw(path: Path) -> pd.DataFrame:
    if path.suffix.lower() == ".csv":
        for enc in ("utf-8-sig", "cp949"):
            try:
                return pd.read_csv(path, header=None, dtype=object, encoding=enc)
            except UnicodeDecodeError:
                continue
        raise ValueError(f"인코딩을 알 수 없음: {path.name}")
    return pd.read_excel(path, header=None, dtype=object)


def locate_header(raw: pd.DataFrame) -> int:
    date_names = {a.lower() for a in ALIASES["date"]}
    for i in range(min(15, len(raw))):
        cells = {str(c).strip().lower() for c in raw.iloc[i].tolist() if pd.notna(c)}
        if cells & date_names:
            return i
    raise ValueError("헤더 행을 찾지 못했습니다")


def map_columns(columns: list[str]) -> dict[str, str]:
    mapping = {}
    for col in columns:
        key = str(col).strip().lower()
        for std, names in ALIASES.items():
            if key in (n.lower() for n in names):
                mapping[col] = std
    missing = set(ALIASES) - set(mapping.values())
    if missing:
        raise ValueError(f"필수 컬럼 누락: {', '.join(sorted(missing))}")
    return mapping


def parse_date(v) -> date | None:
    if isinstance(v, (datetime, pd.Timestamp)):
        return v.date()
    if isinstance(v, date):
        return v
    s = re.sub(r"[^\d]", "", str(v))
    if len(s) == 8:
        try:
            return datetime.strptime(s, "%Y%m%d").date()
        except ValueError:
            return None
    return None


def parse_number(v) -> float | None:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v))
    return float(s) if s not in ("", "-", ".") else None


def load_branch(path: Path, issues: list[dict]) -> pd.DataFrame:
    branch = path.stem.split("_")[0]
    raw = read_raw(path)
    h = locate_header(raw)
    df = raw.iloc[h + 1 :].copy()
    df.columns = [str(c).strip() for c in raw.iloc[h].tolist()]
    df = df.rename(columns=map_columns(list(df.columns)))[list(ALIASES)]
    df["source_row"] = df.index + 1  # 원본 파일 기준 행 번호
    df["branch"] = branch

    def issue(row, reason):
        issues.append({"file": path.name, "row": int(row.source_row), "reason": reason})

    blank = df[list(ALIASES)].isna().all(axis=1) | (df[list(ALIASES)].astype(str).apply(lambda c: c.str.strip()) == "").all(axis=1)
    for _, r in df[blank].iterrows():
        issue(r, "빈 행 제거")
    df = df[~blank].copy()

    df["date"] = df["date"].map(parse_date)
    df["product"] = df["product"].astype(str).str.strip()
    for col in ("qty", "price", "amount"):
        df[col] = df[col].map(parse_number)

    bad = df[["date", "qty", "price", "amount"]].isna().any(axis=1)
    for _, r in df[bad].iterrows():
        issue(r, "형식 오류로 제외 (날짜/숫자 해석 불가)")
    df = df[~bad].copy()

    dup = df.duplicated(subset=["date", "product", "qty", "price", "amount"], keep="first")
    for _, r in df[dup].iterrows():
        issue(r, f"중복 행 제거 — {r['date']} {r['product']}")
    df = df[~dup].copy()

    mismatch = (df["qty"] * df["price"] - df["amount"]).abs() > 1
    for _, r in df[mismatch].iterrows():
        issue(r, f"금액 불일치 — {r['qty']:.0f}×{r['price']:,.0f} ≠ {r['amount']:,.0f} (수량×단가로 보정)")
    df.loc[mismatch, "amount"] = df.loc[mismatch, "qty"] * df.loc[mismatch, "price"]

    for _, r in df[df["qty"] < 0].iterrows():
        issue(r, f"반품 {abs(r['qty']):.0f}건 — 매출에서 차감")

    log.info("%-8s %-28s 헤더 %d행 · 유효 %d행", branch, path.name, h + 1, len(df))
    return df


# ───────────────────────────── 보고서 스타일 ─────────────────────────────

def style_header(ws: Worksheet, row: int, ncols: int, start_col: int = 1) -> None:
    for c in range(start_col, start_col + ncols):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=BONE)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=THIN)


def write_table(ws: Worksheet, df: pd.DataFrame, top: int, left: int = 1, money_cols: tuple[str, ...] = ()) -> tuple[int, int]:
    for j, col in enumerate(df.columns):
        ws.cell(row=top, column=left + j, value=str(col))
    style_header(ws, top, len(df.columns), left)
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            cell = ws.cell(row=top + i, column=left + j, value=v.item() if hasattr(v, "item") else v)
            cell.border = Border(bottom=THIN)
            if df.columns[j] in money_cols:
                cell.number_format = "#,##0"
            if i % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F5F1")
    return top + len(df), left + len(df.columns) - 1


def autosize(ws: Worksheet, min_w: int = 8, max_w: int = 48) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                n = sum(2 if ord(ch) > 127 else 1 for ch in str(c.value))
                widths[c.column] = max(widths.get(c.column, 0), n)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(min_w, min(max_w, w + 2))


def section_title(ws: Worksheet, cell: str, text: str, sub: str = "") -> None:
    ws[cell] = text
    ws[cell].font = Font(bold=True, size=13, color=INK)
    if sub:
        nxt = ws[cell].offset(row=1)
        nxt.value = sub
        nxt.font = Font(size=9, color=ASH, italic=True)


# ───────────────────────────── 보고서 작성 ─────────────────────────────

def build_report(df: pd.DataFrame, issues: list[dict], files: list[Path], out: Path) -> dict:
    df = df.copy()
    df["month"] = pd.to_datetime(df["date"]).dt.strftime("%Y-%m")
    df["weekday"] = pd.to_datetime(df["date"]).dt.day_of_week

    total = int(df["amount"].sum())
    qty = int(df["qty"].sum())
    branch_rev = df.groupby("branch")["amount"].sum().sort_values(ascending=False)
    pivot = df.pivot_table(index="branch", columns="month", values="amount", aggfunc="sum").reindex(branch_rev.index)
    pivot["합계"] = pivot.sum(axis=1)
    products = (
        df.groupby("product").agg(판매수량=("qty", "sum"), 매출=("amount", "sum")).sort_values("매출", ascending=False).reset_index()
    )
    products.insert(0, "순위", range(1, len(products) + 1))
    products["비중(%)"] = (products["매출"] / total * 100).round(1)
    monthly = df.groupby("month")["amount"].sum()
    weekday = df.groupby("weekday")["amount"].sum() / df.groupby("weekday")["date"].nunique()

    wb = Workbook()

    # ── 요약 ──
    ws = wb.active
    ws.title = "요약"
    ws.sheet_view.showGridLines = False
    ws["B2"] = "2026 3분기 매출 통합 보고서"
    ws["B2"].font = Font(bold=True, size=20, color=INK)
    ws["B3"] = f"자동 생성 {datetime.now():%Y-%m-%d %H:%M} · 원본 {len(files)}개 파일 · 지점 {df['branch'].nunique()}곳"
    ws["B3"].font = Font(size=9, color=ASH)

    kpis = [
        ("총 매출", f"₩{total:,}"),
        ("판매 수량", f"{qty:,}개"),
        ("월평균 매출", f"₩{int(monthly.mean()):,}"),
        ("최고 매출 지점", branch_rev.index[0]),
        ("데이터 정제", f"{len(issues)}건 처리"),
    ]
    for i, (label, value) in enumerate(kpis):
        col = 2 + i * 2
        lab = ws.cell(row=5, column=col, value=label)
        val = ws.cell(row=6, column=col, value=value)
        lab.font = Font(size=9, color=ASH)
        val.font = Font(bold=True, size=15, color=VERMILION if i == 0 else INK)
        for r in (5, 6):
            ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=PAPER)
            ws.cell(row=r, column=col + 1).fill = PatternFill("solid", fgColor=PAPER)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 1)
    ws.row_dimensions[6].height = 30

    section_title(ws, "B9", "지점별 매출", "Revenue by Branch")
    br = pd.DataFrame({"지점": branch_rev.index, "매출": branch_rev.values.astype(int)})
    br["비중(%)"] = (br["매출"] / total * 100).round(1)
    end_row, _ = write_table(ws, br, top=11, left=2, money_cols=("매출",))

    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "지점별 매출"
    chart.y_axis.numFmt = "#,##0"
    chart.add_data(Reference(ws, min_col=3, min_row=11, max_row=end_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=2, min_row=12, max_row=end_row))
    chart.series[0].graphicalProperties.solidFill = INK
    chart.x_axis.scaling.orientation = "maxMin"  # 가로 막대: 1위가 위로
    chart.legend = None
    chart.height, chart.width = 7.5, 15
    ws.add_chart(chart, "F9")

    section_title(ws, "B26", "월별 추이", "Monthly Trend")
    mt = pd.DataFrame({"월": monthly.index, "매출": monthly.values.astype(int)})
    m_end, _ = write_table(ws, mt, top=28, left=2, money_cols=("매출",))
    line = LineChart()
    line.title = "월별 매출"
    line.y_axis.numFmt = "#,##0"
    line.add_data(Reference(ws, min_col=3, min_row=28, max_row=m_end), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=2, min_row=29, max_row=m_end))
    line.series[0].graphicalProperties.line.solidFill = VERMILION
    line.series[0].graphicalProperties.line.width = 28000
    line.legend = None
    line.height, line.width = 7.5, 15
    ws.add_chart(line, "F26")
    for col in "BCDEFGHIJK":
        ws.column_dimensions[col].width = 14

    # ── 지점×월 피벗 ──
    ws2 = wb.create_sheet("지점×월")
    ws2.sheet_view.showGridLines = False
    section_title(ws2, "A1", "지점 × 월 매출 피벗", "색이 진할수록 매출이 높습니다")
    pv = pivot.reset_index().rename(columns={"branch": "지점"})
    pv[pv.columns[1:]] = pv[pv.columns[1:]].astype(int)
    p_end, p_right = write_table(ws2, pv, top=4, money_cols=tuple(pv.columns[1:]))
    rng = f"B5:{get_column_letter(p_right - 1)}{p_end}"
    ws2.conditional_formatting.add(rng, ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="F4A595"))
    autosize(ws2, min_w=12)

    # ── 상품 순위 ──
    ws3 = wb.create_sheet("상품 순위")
    ws3.sheet_view.showGridLines = False
    section_title(ws3, "A1", "상품별 매출 순위", "Product Ranking")
    products["판매수량"] = products["판매수량"].astype(int)
    products["매출"] = products["매출"].astype(int)
    write_table(ws3, products, top=4, money_cols=("매출", "판매수량"))
    section_title(ws3, "G1", "요일별 일평균 매출", "Avg. revenue per weekday")
    wd = pd.DataFrame({"요일": [WEEKDAYS[i] for i in weekday.index], "일평균 매출": weekday.values.astype(int)})
    w_end, _ = write_table(ws3, wd, top=4, left=7, money_cols=("일평균 매출",))
    wchart = BarChart()
    wchart.title = "요일별 일평균 매출"
    wchart.add_data(Reference(ws3, min_col=8, min_row=4, max_row=w_end), titles_from_data=True)
    wchart.set_categories(Reference(ws3, min_col=7, min_row=5, max_row=w_end))
    wchart.series[0].graphicalProperties.solidFill = VERMILION
    wchart.legend = None
    wchart.height, wchart.width = 7, 13
    ws3.add_chart(wchart, "J4")
    autosize(ws3, min_w=10)

    # ── 정제 데이터 ──
    ws4 = wb.create_sheet("정제 데이터")
    clean = df[["date", "branch", "product", "qty", "price", "amount"]].sort_values(["date", "branch", "product"])
    clean.columns = ["일자", "지점", "상품", "수량", "단가", "금액"]
    clean[["수량", "단가", "금액"]] = clean[["수량", "단가", "금액"]].astype(int)
    c_end, c_right = write_table(ws4, clean, top=1, money_cols=("단가", "금액"))
    for row in ws4.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "yyyy-mm-dd"
    ws4.freeze_panes = "A2"
    ws4.auto_filter.ref = f"A1:{get_column_letter(c_right)}{c_end}"
    for col, w in zip("ABCDEF", (12, 10, 14, 8, 10, 12)):
        ws4.column_dimensions[col].width = w

    # ── 검증 로그 ──
    ws5 = wb.create_sheet("검증 로그")
    ws5.sheet_view.showGridLines = False
    section_title(ws5, "A1", "데이터 검증 로그", "원본 파일의 몇 번째 행을 어떻게 처리했는지 기록합니다")
    il = pd.DataFrame(issues or [{"file": "-", "row": 0, "reason": "이상 없음"}]).rename(columns={"file": "파일", "row": "원본 행", "reason": "처리 내용"})
    write_table(ws5, il, top=4)
    autosize(ws5)

    for sheet in wb.worksheets:  # 일부 Excel 버전에서 축이 숨겨지는 openpyxl 기본값 보정
        for ch in sheet._charts:
            ch.x_axis.delete = False
            ch.y_axis.delete = False
    wb.save(out)

    kinds: dict[str, int] = {}
    for it in issues:
        key = it["reason"].split(" — ")[0].split(" (")[0]
        key = re.sub(r"\d+건", "", key).strip()
        kinds[key] = kinds.get(key, 0) + 1
    return {
        "files": [f.name for f in files],
        "rows_clean": int(len(df)),
        "issues": len(issues),
        "issues_by_kind": kinds,
        "total_revenue": total,
        "total_qty": qty,
        "branches": [{"name": k, "revenue": int(v)} for k, v in branch_rev.items()],
        "monthly": [{"month": k, "revenue": int(v)} for k, v in monthly.items()],
        "products": [{"name": r["product"], "revenue": int(r["매출"]), "qty": int(r["판매수량"])} for _, r in products.iterrows()],
        "weekday": [{"day": WEEKDAYS[i], "avg": int(v)} for i, v in weekday.items()],
        "sample_issues": issues[:8],
    }


def main() -> None:
    ap = argparse.ArgumentParser(description="지점별 매출 파일 통합 보고서 생성기")
    ap.add_argument("--input", type=Path, default=HERE / "input")
    ap.add_argument("--out", type=Path, default=HERE / "output")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[logging.StreamHandler(), logging.FileHandler(args.out / "run.log", "w", encoding="utf-8")],
    )

    started = time.perf_counter()
    files = sorted(p for p in args.input.iterdir() if p.suffix.lower() in (".xlsx", ".xls", ".csv") and not p.name.startswith("~$"))
    if not files:
        raise SystemExit(f"{args.input} 에 입력 파일이 없습니다. 먼저 make_sample_data.py 를 실행하세요.")

    issues: list[dict] = []
    frames = []
    for f in files:
        try:
            frames.append(load_branch(f, issues))
        except Exception as exc:
            issues.append({"file": f.name, "row": 0, "reason": f"파일 전체 제외 — {exc}"})
            log.error("%s 처리 실패: %s", f.name, exc)
    df = pd.concat(frames, ignore_index=True)
    issues.sort(key=lambda it: (it["file"], it["row"]))

    report = args.out / "매출_통합보고서.xlsx"
    summary = build_report(df, issues, files, report)
    summary["elapsed_sec"] = round(time.perf_counter() - started, 2)
    (args.out / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    log.info("완료 — %d개 파일 · 유효 %d행 · 정제 %d건 · 총매출 ₩%s · %.1fs → %s",
             len(files), summary["rows_clean"], summary["issues"], f"{summary['total_revenue']:,}", summary["elapsed_sec"], report)


if __name__ == "__main__":
    main()
