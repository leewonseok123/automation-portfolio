"""
05 · AI Document Extractor — PDF 문서(세금계산서·견적서·영수증) → 검증된 엑셀

흐름
  PDF 글자 추출 → LLM(DeepSeek) 이 JSON 으로 구조화 → 정규화 → 자동 검증 → 엑셀·JSON 저장

설계 원칙
- AI 가 계산해서 채우지 않게: 문서에 '인쇄된 값'만 옮기게 하고, 합계·세액·사업자번호는 코드로 따로 검증합니다.
  그래야 원본 문서의 오타(합계 불일치 등)를 AI 가 슬쩍 고쳐버리지 않고 '검토 필요'로 드러냅니다.
- 글자 정보가 없는 스캔본은 API 를 호출하지 않고 'OCR 필요' 로 표시 (비용 절약, 환각 방지)
- 모든 결과에 상태(ok / review / needs_ocr / skipped / failed)와 사유를 남김
- 정답지(ground_truth.json)가 있으면 필드별 정확도를 계산 → 규칙 기반(rules)과 AI 비교

사용법
    python ai_extract.py --provider rules                 # 키 없이: 규칙 기반 기준선
    python ai_extract.py --provider deepseek --dry-run    # API 호출 없이 대상·분량만 확인
    python ai_extract.py --provider deepseek              # 실제 AI 추출 (DEEPSEEK_API_KEY 필요)

환경변수 (또는 프로젝트 루트의 .env)
    DEEPSEEK_API_KEY    필수
    DEEPSEEK_MODEL      기본 deepseek-chat
    DEEPSEEK_BASE_URL   기본 https://api.deepseek.com
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import re
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Protocol

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

HERE = Path(__file__).parent
ROOT = HERE.parent.parent
log = logging.getLogger("extract")

TYPE_KO = {"tax_invoice": "세금계산서", "quotation": "견적서", "receipt": "영수증", "other": "기타"}
MAX_CHARS = 12_000  # 문서 1건당 LLM 에 보내는 최대 글자 수
MIN_TEXT_CHARS = 20  # 이보다 적으면 스캔본으로 간주

SYSTEM_PROMPT = """너는 한국 사업장 문서(세금계산서·견적서·영수증)에서 정보를 옮겨 적는 도구다.

규칙
1. 문서에 인쇄된 값만 그대로 옮긴다. 추측하거나 계산해서 채우지 않는다. 합계가 안 맞아 보여도 인쇄된 값을 그대로 적는다.
2. 문서에 없는 항목은 null.
3. 금액은 원 단위 정수, 날짜는 YYYY-MM-DD, 사업자등록번호는 000-00-00000 형식.
4. 공급자는 물건·서비스를 판 쪽, 공급받는자(수신)는 산 쪽이다. 영수증은 공급자만 있고 buyer 는 null.
5. 세금계산서·견적서·영수증이 아니면 doc_type 을 "other" 로 하고 나머지는 null/빈 배열.
6. 설명 없이 JSON 객체 하나만 출력한다.

출력 형식
{"doc_type": "tax_invoice|quotation|receipt|other",
 "issue_date": "YYYY-MM-DD",
 "supplier": {"name": "", "biz_no": ""},
 "buyer": {"name": ""} 또는 null,
 "items": [{"name": "", "qty": 0, "unit_price": 0, "amount": 0}],
 "supply_amount": 0, "vat": 0, "total": 0}"""


# ───────────────────────────── 환경 · 데이터 모델 ─────────────────────────────


def load_env() -> None:
    """.env 를 읽어 환경변수로 (이미 설정된 값은 덮어쓰지 않음). 키 값은 절대 출력하지 않습니다."""
    for path in (ROOT / ".env", HERE / ".env"):
        if not path.exists():
            continue
        for ln in path.read_text(encoding="utf-8").splitlines():
            ln = ln.strip()
            if ln and not ln.startswith("#") and "=" in ln:
                k, v = ln.split("=", 1)
                os.environ.setdefault(k.strip(), v.strip().strip("'\""))


@dataclass
class Usage:
    prompt: int = 0
    completion: int = 0

    def add(self, other: Usage) -> None:
        self.prompt += other.prompt
        self.completion += other.completion


@dataclass
class Issue:
    code: str
    severity: str  # error | warn
    message: str


@dataclass
class Result:
    file: str
    status: str  # ok | review | needs_ocr | skipped | failed
    doc: dict | None = None
    issues: list[Issue] = field(default_factory=list)
    usage: Usage = field(default_factory=Usage)
    error: str | None = None


# ───────────────────────────── PDF → 글자 ─────────────────────────────


def read_pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in reader.pages).strip()


# ───────────────────────────── 추출기 ─────────────────────────────


class Extractor(Protocol):
    name: str
    model: str

    def extract(self, text: str) -> tuple[dict, Usage]: ...


def parse_json(raw: str) -> dict:
    """코드블록(```json)이나 앞뒤 설명이 붙어 있어도 JSON 객체를 꺼냅니다."""
    s = raw.strip()
    s = re.sub(r"^```(?:json)?\s*|\s*```$", "", s, flags=re.IGNORECASE)
    start, end = s.find("{"), s.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("JSON 객체를 찾지 못했습니다")
    obj = json.loads(s[start : end + 1])
    if not isinstance(obj, dict):
        raise ValueError("JSON 최상위가 객체가 아닙니다")
    return obj


class DeepSeekExtractor:
    name = "deepseek"

    def __init__(self, attempts: int = 3):
        from openai import OpenAI  # DeepSeek API 는 OpenAI SDK 와 호환

        key = os.getenv("DEEPSEEK_API_KEY")
        if not key:
            raise SystemExit("DEEPSEEK_API_KEY 가 없습니다. 프로젝트 루트의 .env 에 넣거나 환경변수로 설정하세요.")
        self.model = os.getenv("DEEPSEEK_MODEL", "deepseek-chat")
        self.client = OpenAI(
            api_key=key, base_url=os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com"), timeout=90, max_retries=2
        )
        self.attempts = attempts

    def extract(self, text: str) -> tuple[dict, Usage]:
        usage = Usage()
        messages = [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f'문서 텍스트:\n"""\n{text[:MAX_CHARS]}\n"""'},
        ]
        last_err = "알 수 없는 오류"
        for attempt in range(1, self.attempts + 1):
            resp = self.client.chat.completions.create(
                model=self.model, messages=messages, temperature=0, response_format={"type": "json_object"}
            )
            if resp.usage:
                usage.add(Usage(resp.usage.prompt_tokens, resp.usage.completion_tokens))
            raw = resp.choices[0].message.content or ""
            try:
                return parse_json(raw), usage
            except ValueError as exc:
                last_err = str(exc)
                log.warning("JSON 해석 실패 (%d/%d): %s", attempt, self.attempts, exc)
                messages += [
                    {"role": "assistant", "content": raw},
                    {"role": "user", "content": "올바른 JSON 객체가 아닙니다. 설명 없이 JSON 객체 하나만 다시 출력하세요."},
                ]
        raise ValueError(f"JSON 형식 응답을 받지 못했습니다: {last_err}")


class RulesExtractor:
    """기준선: 샘플 한 장을 보고 개발자가 흔히 짜는 정규식. 양식이 달라지면 깨지는 것을 보여주기 위한 비교 대상입니다."""

    name = "rules"
    model = "regex"

    def extract(self, text: str) -> tuple[dict, Usage]:
        def num(label: str) -> int | None:
            m = re.search(label + r"\s*[:：]?\s*([\d,]+)\s*$", text, re.MULTILINE)
            return int(m.group(1).replace(",", "")) if m else None

        if "세금계산서" in text:
            kind = "tax_invoice"
        elif "견 적 서" in text or "견적서" in text:
            kind = "quotation"
        elif "카드결제" in text:
            kind = "receipt"
        else:
            kind = "other"
        if kind == "other":
            return {"doc_type": "other", "items": []}, Usage()
        date_m = re.search(r"(?:작성일자|견적일)\s*[:：]\s*(\d{4}-\d{2}-\d{2})", text)
        sup = re.search(r"\[공급자\]\s*상호:\s*(\S+)", text) or re.search(r"공급자 상호:\s*(\S+)", text)
        bn = re.search(r"사업자등록번호:\s*(\d{3}-\d{2}-\d{5})", text)
        buyer = re.search(r"\[공급받는자\]\s*상호:\s*(\S+)", text) or re.search(r"수신:\s*(\S+)", text)
        return {
            "doc_type": kind,
            "issue_date": date_m.group(1) if date_m else None,
            "supplier": {"name": sup.group(1) if sup else None, "biz_no": bn.group(1) if bn else None},
            "buyer": {"name": buyer.group(1)} if buyer else None,
            "items": [],
            "supply_amount": num("(?:공급가액 합계|소계)"),
            "vat": num("(?:세액|부가가치세\\(10%\\))"),
            "total": num("(?:합계금액|총 견적금액)"),
        }, Usage()


# ───────────────────────────── 정규화 ─────────────────────────────


def to_int(v) -> int | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    digits = re.sub(r"[^\d-]", "", str(v))
    return int(digits) if digits not in ("", "-") else None


def to_iso_date(v) -> str | None:
    if not v:
        return None
    s = str(v).strip()
    m = re.search(r"(\d{4})\D+(\d{1,2})\D+(\d{1,2})", s) or re.search(r"\b(\d{2})\D(\d{1,2})\D(\d{1,2})\b", s)
    if not m:
        return None
    y, mo, d = (int(x) for x in m.groups())
    y = y + 2000 if y < 100 else y
    try:
        return date(y, mo, d).isoformat()
    except ValueError:
        return None


def to_biz_no(v) -> str | None:
    if not v:
        return None
    digits = re.sub(r"\D", "", str(v))
    return f"{digits[:3]}-{digits[3:5]}-{digits[5:]}" if len(digits) == 10 else str(v).strip()


def normalize(raw: dict) -> dict:
    kind = str(raw.get("doc_type") or "other").strip().lower()
    kind = kind if kind in TYPE_KO else "other"
    sup = raw.get("supplier") if isinstance(raw.get("supplier"), dict) else None
    buyer = raw.get("buyer") if isinstance(raw.get("buyer"), dict) else None
    items = []
    for it in raw.get("items") or []:
        if isinstance(it, dict):
            items.append(
                {
                    "name": str(it.get("name") or "").strip(),
                    "qty": to_int(it.get("qty")),
                    "unit_price": to_int(it.get("unit_price")),
                    "amount": to_int(it.get("amount")),
                }
            )
    raw_date = raw.get("issue_date")
    return {
        "doc_type": kind,
        "issue_date": to_iso_date(raw_date),
        "issue_date_raw": raw_date if raw_date and not to_iso_date(raw_date) else None,
        "supplier": {"name": (sup.get("name") or "").strip() or None, "biz_no": to_biz_no(sup.get("biz_no"))} if sup else None,
        "buyer": {"name": (buyer.get("name") or "").strip() or None} if buyer else None,
        "items": items,
        "supply_amount": to_int(raw.get("supply_amount")),
        "vat": to_int(raw.get("vat")),
        "total": to_int(raw.get("total")),
    }


# ───────────────────────────── 검증 ─────────────────────────────


def biz_no_valid(biz_no: str | None) -> bool:
    """사업자등록번호 검증 숫자 확인 (가중치 1·3·7·1·3·7·1·3·5)."""
    digits = re.sub(r"\D", "", biz_no or "")
    if len(digits) != 10:
        return False
    d = [int(c) for c in digits]
    w = [1, 3, 7, 1, 3, 7, 1, 3, 5]
    s = sum(a * b for a, b in zip(d[:9], w, strict=True)) + (d[8] * 5) // 10
    return (10 - s % 10) % 10 == d[9]


def validate(doc: dict) -> list[Issue]:
    if doc["doc_type"] == "other":
        return []
    out: list[Issue] = []
    kind, supply, vat, total = doc["doc_type"], doc["supply_amount"], doc["vat"], doc["total"]

    for key, label in (("issue_date", "일자"), ("total", "합계금액")):
        if doc[key] is None:
            out.append(Issue("MISSING", "warn", f"{label}를 찾지 못했습니다"))
    if not doc["supplier"] or not doc["supplier"]["name"]:
        out.append(Issue("MISSING", "warn", "공급자 상호를 찾지 못했습니다"))
    if doc.get("issue_date_raw"):
        out.append(Issue("BAD_DATE", "warn", f"날짜를 해석할 수 없습니다: {doc['issue_date_raw']}"))

    bn = doc["supplier"]["biz_no"] if doc["supplier"] else None
    if not bn:
        out.append(Issue("MISSING", "warn", "공급자 사업자번호를 찾지 못했습니다"))
    elif not biz_no_valid(bn):
        out.append(Issue("BIZNO_CHECKSUM", "error", f"사업자등록번호 {bn} 의 검증 숫자가 맞지 않습니다 (오타 가능성)"))

    items_sum = sum(i["amount"] or 0 for i in doc["items"])
    base = total if kind == "receipt" else supply  # 영수증 품목은 부가세 포함가
    if doc["items"] and base is not None and abs(items_sum - base) > 1:
        out.append(
            Issue(
                "ITEMS_SUM", "error", f"품목 금액 합계 {items_sum:,} ≠ {'합계금액' if kind == 'receipt' else '공급가액'} {base:,}"
            )
        )

    if vat is not None and supply is not None:
        expect = round(supply * 0.1)
        if kind != "receipt" and abs(vat - expect) > 1:
            out.append(Issue("VAT_RATE", "error", f"세액 {vat:,} ≠ 공급가액의 10% ({expect:,})"))
        if abs(supply + vat - (total or 0)) > 1 and total is not None:
            out.append(Issue("TOTAL_SUM", "error", f"공급가액 {supply:,} + 세액 {vat:,} ≠ 합계금액 {total:,}"))
    return out


# ───────────────────────────── 처리 · 평가 ─────────────────────────────


def process(path: Path, extractor: Extractor) -> Result:
    try:
        text = read_pdf_text(path)
    except Exception as exc:
        return Result(path.name, "failed", error=f"PDF 를 읽을 수 없습니다: {exc}")
    if len(text) < MIN_TEXT_CHARS:
        issue = Issue("NEEDS_OCR", "warn", "글자 정보가 없는 스캔본입니다. OCR 처리 후 다시 시도하세요 (AI 호출 안 함)")
        return Result(path.name, "needs_ocr", issues=[issue])
    try:
        raw, usage = extractor.extract(text)
        doc = normalize(raw)
    except Exception as exc:
        return Result(path.name, "failed", error=str(exc)[:300])
    if doc["doc_type"] == "other":
        return Result(
            path.name,
            "skipped",
            doc=doc,
            usage=usage,
            issues=[Issue("NOT_INVOICE", "warn", "세금계산서·견적서·영수증이 아닌 문서")],
        )
    issues = validate(doc)
    return Result(path.name, "review" if issues else "ok", doc=doc, issues=issues, usage=usage)


EVAL_FIELDS = [
    "doc_type",
    "issue_date",
    "supplier_name",
    "supplier_biz_no",
    "buyer_name",
    "supply_amount",
    "vat",
    "total",
    "items",
]


def flatten(doc: dict | None) -> dict:
    d = doc or {}
    sup, buyer = d.get("supplier") or {}, d.get("buyer") or {}
    return {
        "doc_type": d.get("doc_type"),
        "issue_date": d.get("issue_date"),
        "supplier_name": sup.get("name"),
        "supplier_biz_no": sup.get("biz_no"),
        "buyer_name": buyer.get("name"),
        "supply_amount": d.get("supply_amount"),
        "vat": d.get("vat"),
        "total": d.get("total"),
        "items": sorted((i.get("amount") or 0) for i in d.get("items") or []),
    }


def evaluate(results: list[Result], truth: dict[str, dict]) -> dict:
    """OCR 이 필요한 문서는 제외하고, 읽을 수 있었던 문서의 필드별 정확도를 계산합니다."""
    per_field = {f: [0, 0] for f in EVAL_FIELDS}
    docs_perfect = docs_scored = 0
    per_doc = {}
    for r in results:
        t = truth.get(r.file)
        if t is None or r.status == "needs_ocr":
            continue
        want, got = flatten({**t, "items": t.get("items", [])}), flatten(r.doc)
        fields = ["doc_type"] if want["doc_type"] == "other" else EVAL_FIELDS  # 기타 문서는 분류만 평가
        wrong = []
        for f in fields:
            per_field[f][1] += 1
            if want[f] == got[f]:
                per_field[f][0] += 1
            else:
                wrong.append(f)
        docs_scored += 1
        docs_perfect += not wrong
        per_doc[r.file] = wrong
    correct = sum(c for c, _ in per_field.values())
    total = sum(n for _, n in per_field.values())
    return {
        "docs_scored": docs_scored,
        "docs_perfect": docs_perfect,
        "field_accuracy": round(correct / total * 100, 1) if total else None,
        "by_field": {f: {"correct": c, "total": n} for f, (c, n) in per_field.items() if n},
        "wrong_fields": per_doc,
    }


# ───────────────────────────── 엑셀 ─────────────────────────────

INK, BONE, PAPER, VERMILION = "131311", "F2EFE9", "E9E5DD", "E8442E"
STATUS_FILL = {"ok": "E3F4EA", "review": "FBE1DC", "needs_ocr": "ECECEC", "skipped": "ECECEC", "failed": "F6C7C0"}
STATUS_KO = {"ok": "정상", "review": "검토 필요", "needs_ocr": "OCR 필요", "skipped": "대상 아님", "failed": "실패"}


def _header(ws, row: int, cols: list[str]) -> None:
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font = Font(bold=True, color=BONE)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center")


def _autosize(ws, cap: int = 60) -> None:
    for col in ws.columns:
        width = max((sum(2 if ord(ch) > 127 else 1 for ch in str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(cap, width + 2)


def write_xlsx(results: list[Result], summary: dict, path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "추출결과"
    _header(
        ws, 1, ["파일", "종류", "일자", "공급자", "사업자번호", "공급받는자", "공급가액", "세액", "합계금액", "상태", "검토 사항"]
    )
    for r in results:
        d = r.doc or {}
        sup, buyer = d.get("supplier") or {}, d.get("buyer") or {}
        notes = " / ".join(i.message for i in r.issues) or (r.error or "")
        ws.append(
            [
                r.file,
                TYPE_KO.get(d.get("doc_type"), "-"),
                d.get("issue_date"),
                sup.get("name"),
                sup.get("biz_no"),
                buyer.get("name"),
                d.get("supply_amount"),
                d.get("vat"),
                d.get("total"),
                STATUS_KO[r.status],
                notes,
            ]
        )
        fill = PatternFill("solid", fgColor=STATUS_FILL[r.status])
        for cell in ws[ws.max_row]:
            cell.fill = fill
        for col in (7, 8, 9):
            ws.cell(row=ws.max_row, column=col).number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    _autosize(ws)

    wi = wb.create_sheet("품목")
    _header(wi, 1, ["파일", "품목", "수량", "단가", "금액"])
    for r in results:
        for it in (r.doc or {}).get("items", []):
            wi.append([r.file, it["name"], it["qty"], it["unit_price"], it["amount"]])
            for col in (4, 5):
                wi.cell(row=wi.max_row, column=col).number_format = "#,##0"
    wi.freeze_panes = "A2"
    _autosize(wi)

    wr = wb.create_sheet("검토 필요")
    _header(wr, 1, ["파일", "심각도", "코드", "내용"])
    for r in results:
        for i in r.issues:
            wr.append([r.file, "오류" if i.severity == "error" else "확인", i.code, i.message])
            if i.severity == "error":
                wr.cell(row=wr.max_row, column=2).font = Font(bold=True, color=VERMILION)
    if wr.max_row == 1:
        wr.append(["-", "-", "-", "검토할 항목이 없습니다"])
    _autosize(wr)

    wsum = wb.create_sheet("요약")
    wsum.sheet_view.showGridLines = False
    wsum["A1"] = "문서 추출 요약"
    wsum["A1"].font = Font(bold=True, size=16)
    rows = [("추출 방식", f"{summary['provider']} ({summary['model']})"), ("처리 문서", summary["docs"])]
    rows += [(STATUS_KO[k], v) for k, v in summary["status_counts"].items()]
    rows += [("토큰(입력/출력)", f"{summary['tokens']['prompt']:,} / {summary['tokens']['completion']:,}")]
    if summary.get("accuracy"):
        a = summary["accuracy"]
        rows += [("필드 정확도", f"{a['field_accuracy']}%"), ("완전 일치 문서", f"{a['docs_perfect']} / {a['docs_scored']}")]
    for i, (k, v) in enumerate(rows, start=3):
        wsum.cell(row=i, column=1, value=k).font = Font(bold=True)
        wsum.cell(row=i, column=2, value=v)
    _autosize(wsum)
    wb.save(path)


# ───────────────────────────── 실행 ─────────────────────────────


def main() -> None:
    load_env()
    ap = argparse.ArgumentParser(description="PDF 문서 → 검증된 엑셀")
    ap.add_argument("--provider", choices=["deepseek", "rules"], default="rules")
    ap.add_argument("--input", type=Path, default=HERE / "input")
    ap.add_argument("--out", type=Path, default=None, help="기본: output/<provider>")
    ap.add_argument("--workers", type=int, default=4)
    ap.add_argument("--limit", type=int, default=50, help="한 번에 처리할 최대 문서 수 (비용 안전장치)")
    ap.add_argument("--dry-run", action="store_true", help="API 호출 없이 대상 문서와 분량만 확인")
    args = ap.parse_args()

    out = args.out or HERE / "output" / args.provider
    out.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[logging.StreamHandler(), logging.FileHandler(out / "run.log", "w", encoding="utf-8")],
    )
    files = sorted(args.input.glob("*.pdf"))[: args.limit]
    if not files:
        raise SystemExit(f"{args.input} 에 PDF 가 없습니다. 먼저 make_sample_docs.py 를 실행하세요.")

    if args.dry_run:
        chars = {f.name: len(read_pdf_text(f)) for f in files}
        readable = {k: v for k, v in chars.items() if v >= MIN_TEXT_CHARS}
        log.info(
            "대상 %d건 / AI 호출 대상 %d건 / OCR 필요 %d건 / 입력 글자 합계 %s자 (API 호출 없음)",
            len(files),
            len(readable),
            len(files) - len(readable),
            f"{sum(readable.values()):,}",
        )
        return

    extractor: Extractor = DeepSeekExtractor() if args.provider == "deepseek" else RulesExtractor()
    started = time.perf_counter()
    log.info("시작 — %d건, 방식: %s (%s)", len(files), extractor.name, extractor.model)
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        results = list(pool.map(lambda f: process(f, extractor), files))
    for r in results:
        detail = ", ".join(i.code for i in r.issues) or (r.error or "")
        log.info("%-24s %-9s %s", r.file, STATUS_KO[r.status], detail)

    truth_path = args.input / "ground_truth.json"
    truth = json.loads(truth_path.read_text(encoding="utf-8")) if truth_path.exists() else None
    tokens = Usage()
    for r in results:
        tokens.add(r.usage)
    summary = {
        "provider": extractor.name,
        "model": extractor.model,
        "docs": len(results),
        "status_counts": dict(Counter(r.status for r in results)),
        "issues_by_code": dict(Counter(i.code for r in results for i in r.issues)),
        "tokens": {"prompt": tokens.prompt, "completion": tokens.completion},
        "elapsed_sec": round(time.perf_counter() - started, 2),
        "accuracy": evaluate(results, truth) if truth else None,
        "results": [
            {"file": r.file, "status": r.status, "doc": r.doc, "issues": [vars(i) for i in r.issues], "error": r.error}
            for r in results
        ],
    }
    write_xlsx(results, summary, out / "extracted.xlsx")
    (out / "extracted.json").write_text(json.dumps(summary["results"], ensure_ascii=False, indent=2), encoding="utf-8")
    (out / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    acc = summary["accuracy"]
    log.info(
        "완료 — %s%s / %.1fs → %s",
        summary["status_counts"],
        f" / 필드 정확도 {acc['field_accuracy']}% (완전 일치 {acc['docs_perfect']}/{acc['docs_scored']})" if acc else "",
        summary["elapsed_sec"],
        out,
    )


if __name__ == "__main__":
    main()
