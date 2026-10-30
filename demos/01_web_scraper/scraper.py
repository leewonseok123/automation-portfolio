"""
01 · Web Scraper — 목록 + 상세 페이지 2단계 수집기

대상: https://books.toscrape.com (스크래핑 연습용으로 공개된 샌드박스 사이트)

기능
- robots.txt 확인 후 허용된 경로만 수집
- User-Agent 로테이션, 지수 백오프 재시도, 요청 간 랜덤 딜레이
- 상세 페이지 병렬 수집 (ThreadPoolExecutor, 동시성 제한)
- CSV / JSON / XLSX 동시 저장 + 실행 요약(run_summary.json)

사용법
    python scraper.py --pages 3 --workers 6 --out output
"""

from __future__ import annotations

import argparse
import csv
import json
import logging
import random
import re
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, fields
from pathlib import Path
from urllib.parse import urljoin
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/"
LIST_URL = BASE_URL + "catalogue/page-{page}.html"

USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/130.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_6) AppleWebKit/605.1.15 (KHTML, like Gecko) Version/18.0 Safari/605.1.15",
    "Mozilla/5.0 (X11; Linux x86_64; rv:131.0) Gecko/20100101 Firefox/131.0",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:131.0) Gecko/20100101 Firefox/131.0",
]

RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}

log = logging.getLogger("scraper")


@dataclass
class Book:
    upc: str
    title: str
    category: str
    price_gbp: float
    rating: int
    stock: int
    url: str
    image_url: str


class PoliteSession:
    """재시도·딜레이·UA 로테이션을 한곳에서 처리하는 얇은 래퍼."""

    def __init__(self, min_delay: float = 0.3, max_delay: float = 0.9, retries: int = 4):
        self.session = requests.Session()
        self.min_delay = min_delay
        self.max_delay = max_delay
        self.retries = retries
        self.request_count = 0
        self.retry_count = 0

    def get(self, url: str) -> str:
        for attempt in range(1, self.retries + 1):
            time.sleep(random.uniform(self.min_delay, self.max_delay))
            try:
                resp = self.session.get(
                    url,
                    headers={"User-Agent": random.choice(USER_AGENTS)},
                    timeout=15,
                )
                self.request_count += 1
                if resp.status_code == 429 or resp.status_code >= 500:
                    raise requests.HTTPError(f"HTTP {resp.status_code}")
                resp.raise_for_status()
                resp.encoding = "utf-8"
                return resp.text
            except requests.RequestException as exc:
                if attempt == self.retries:
                    raise
                self.retry_count += 1
                backoff = 2**attempt + random.random()
                log.warning("retry %d/%d %s (%s) — %.1fs 대기", attempt, self.retries, url, exc, backoff)
                time.sleep(backoff)
        raise RuntimeError("unreachable")


def check_robots(paths: list[str]) -> None:
    rp = RobotFileParser(urljoin(BASE_URL, "robots.txt"))
    try:
        rp.read()
    except Exception:  # robots.txt 가 없으면 기본 허용
        return
    for path in paths:
        if not rp.can_fetch("*", path):
            raise SystemExit(f"robots.txt 가 {path} 수집을 허용하지 않습니다.")


def parse_list(html: str, page_url: str) -> list[str]:
    soup = BeautifulSoup(html, "lxml")
    return [urljoin(page_url, a["href"]) for a in soup.select("article.product_pod h3 a")]


def parse_detail(html: str, url: str) -> Book:
    soup = BeautifulSoup(html, "lxml")
    table = {row.th.get_text(strip=True): row.td.get_text(strip=True) for row in soup.select("table.table tr")}
    rating_cls = soup.select_one("div.product_main p.star-rating")["class"]
    stock_match = re.search(r"\((\d+) available\)", table.get("Availability", ""))
    return Book(
        upc=table["UPC"],
        title=soup.select_one("div.product_main h1").get_text(strip=True),
        category=soup.select("ul.breadcrumb li a")[-1].get_text(strip=True),
        price_gbp=float(re.sub(r"[^\d.]", "", table["Price (incl. tax)"])),
        rating=next((RATING_WORDS[c] for c in rating_cls if c in RATING_WORDS), 0),
        stock=int(stock_match.group(1)) if stock_match else 0,
        url=url,
        image_url=urljoin(url, soup.select_one("#product_gallery img")["src"]),
    )


def save_csv(books: list[Book], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: 엑셀에서 한글 깨짐 방지
        writer = csv.DictWriter(f, fieldnames=[fl.name for fl in fields(Book)])
        writer.writeheader()
        writer.writerows(asdict(b) for b in books)


def save_xlsx(books: list[Book], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "books"
    headers = [fl.name for fl in fields(Book)]
    ws.append(headers)
    for b in books:
        ws.append(list(asdict(b).values()))
    head_fill = PatternFill("solid", fgColor="131311")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="F2EFE9")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for idx, header in enumerate(headers, start=1):
        width = max(len(header), *(len(str(getattr(b, header))) for b in books)) if books else len(header)
        ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser(description="books.toscrape.com 2단계 수집기")
    ap.add_argument("--pages", type=int, default=3, help="수집할 목록 페이지 수 (1페이지 = 20권)")
    ap.add_argument("--workers", type=int, default=6, help="상세 페이지 동시 요청 수")
    ap.add_argument("--out", type=Path, default=Path(__file__).parent / "output")
    args = ap.parse_args()

    args.out.mkdir(parents=True, exist_ok=True)
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)-7s %(message)s",
        datefmt="%H:%M:%S",
        handlers=[logging.StreamHandler(), logging.FileHandler(args.out / "run.log", "w", encoding="utf-8")],
    )
    started = time.perf_counter()

    check_robots(["/catalogue/"])
    client = PoliteSession()

    detail_urls: list[str] = []
    for page in range(1, args.pages + 1):
        url = LIST_URL.format(page=page)
        links = parse_list(client.get(url), url)
        detail_urls.extend(links)
        log.info("목록 %d/%d 페이지 — 링크 %d개", page, args.pages, len(links))

    books: list[Book] = []
    failed: list[str] = []
    with ThreadPoolExecutor(max_workers=args.workers) as pool:
        futures = {pool.submit(client.get, u): u for u in detail_urls}
        for i, fut in enumerate(as_completed(futures), start=1):
            url = futures[fut]
            try:
                books.append(parse_detail(fut.result(), url))
            except Exception as exc:
                failed.append(url)
                log.error("실패 %s: %s", url, exc)
            if i % 10 == 0 or i == len(futures):
                log.info("상세 %d/%d 수집", i, len(futures))

    books.sort(key=lambda b: (b.category, b.title))
    save_csv(books, args.out / "books.csv")
    save_xlsx(books, args.out / "books.xlsx")
    (args.out / "books.json").write_text(json.dumps([asdict(b) for b in books], ensure_ascii=False, indent=2), encoding="utf-8")

    elapsed = time.perf_counter() - started
    summary = {
        "source": BASE_URL,
        "pages": args.pages,
        "items": len(books),
        "failed": len(failed),
        "requests": client.request_count,
        "retries": client.retry_count,
        "workers": args.workers,
        "elapsed_sec": round(elapsed, 2),
        "categories": len({b.category for b in books}),
        "avg_price_gbp": round(sum(b.price_gbp for b in books) / len(books), 2) if books else 0,
        "finished_at": time.strftime("%Y-%m-%dT%H:%M:%S%z"),
    }
    (args.out / "run_summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    log.info("완료 — %d건 / 실패 %d / 요청 %d / %.1fs → %s", len(books), len(failed), client.request_count, elapsed, args.out)


if __name__ == "__main__":
    main()
